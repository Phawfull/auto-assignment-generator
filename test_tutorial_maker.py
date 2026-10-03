import copy
import json
import os
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from docx import Document

import uuid

import ai_answers
import java_programs
import settings
import word_document
import workflow


CONSOLE = {
    "main_class": "SumExample", "mode": "console",
    "files": [{"name": "SumExample.java", "code": """import java.util.Scanner;
public class SumExample {
    public static void main(String[] args) {
        Scanner input = new Scanner(System.in);
        int first = input.nextInt();
        int second = input.nextInt();
        System.out.println("Inputs: " + first + " and " + second);
        System.out.println("Sum: " + (first + second));
    }
}
"""}], "sql": [], "stdin": "17\n25\n", "assumptions": [],
    "steps": [
        {"action": "assert", "target": "consoleOutput", "value": "Sum: 42"},
        {"action": "capture", "target": "", "value": "Adding 17 and 25"},
    ],
}

SWING = {
    "main_class": "GradeExample", "mode": "swing",
    "files": [{"name": "GradeExample.java", "code": """import javax.swing.*;
import java.awt.*;
public class GradeExample {
    public static void main(String[] args) {
        SwingUtilities.invokeLater(() -> {
            JFrame frame = new JFrame("Grade calculator test");
            JPanel panel = new JPanel(new GridLayout(5, 2, 10, 10));
            panel.setBorder(BorderFactory.createEmptyBorder(20, 20, 20, 20));
            JTextField marks = new JTextField();
            marks.setName("marks");
            JLabel result = new JLabel("Enter marks");
            result.setName("result");
            JButton button = new JButton("Calculate");
            button.setName("calculate");
            button.addActionListener(event -> {
                int value = Integer.parseInt(marks.getText());
                result.setText(value >= 90 ? "Grade: A" : "Grade: B");
            });
            JButton popup = new JButton("Show message");
            popup.setName("popup");
            popup.addActionListener(event -> JOptionPane.showMessageDialog(frame, result.getText()));
            panel.add(new JLabel("Marks"));
            panel.add(marks);
            panel.add(button);
            panel.add(result);
            panel.add(popup);
            frame.add(panel);
            frame.setSize(540, 300);
            frame.setLocationRelativeTo(null);
            frame.setVisible(true);
        });
    }
}
"""}], "sql": [], "stdin": "", "assumptions": [],
    "steps": [
        {"action": "text", "target": "marks", "value": "95"},
        {"action": "click", "target": "calculate", "value": ""},
        {"action": "assert", "target": "result", "value": "Grade: A"},
        {"action": "capture", "target": "", "value": "Calculated grade for 95 marks"},
        {"action": "click", "target": "popup", "value": ""},
        {"action": "assert", "target": "", "value": "Grade: A"},
        {"action": "capture", "target": "", "value": "Result dialog"},
        {"action": "click", "target": "OK", "value": ""},
    ],
}


class CoreTests(unittest.TestCase):
    def test_failed_run_is_repaired_before_returning(self):
        from unittest.mock import Mock

        ai = Mock()
        ai.ask.side_effect = [copy.deepcopy(CONSOLE), copy.deepcopy(CONSOLE)]
        options = {"key": "test-key", "db_password": ""}
        with tempfile.TemporaryDirectory(dir=settings.ROOT) as directory:
            with patch("workflow.compile_and_run", side_effect=[RuntimeError("wrong result"), []]):
                result = workflow.solve_question(
                    "Add two numbers", 1, Path(directory), options,
                    ai, threading.Event(), lambda message: None,
                )
        self.assertNotIn("error", result)
        self.assertEqual(ai.ask.call_count, 2)
        self.assertIn("wrong result", ai.ask.call_args.args[1])

    def test_failed_question_keeps_error_after_three_attempts(self):
        from unittest.mock import Mock

        ai = Mock()
        ai.ask.side_effect = [copy.deepcopy(CONSOLE) for _ in range(3)]
        options = {"key": "test-key", "db_password": "private-password"}
        with tempfile.TemporaryDirectory(dir=settings.ROOT) as directory:
            with patch("workflow.compile_and_run", side_effect=RuntimeError("private-password failed")):
                result = workflow.solve_question(
                    "Add two numbers", 1, Path(directory), options,
                    ai, threading.Event(), lambda message: None,
                )
        self.assertEqual(result["error"], "[redacted] failed")
        self.assertEqual(ai.ask.call_count, 3)

    def test_preserves_table_in_question_order(self):
        with tempfile.TemporaryDirectory(dir=settings.ROOT) as directory:
            path = Path(directory) / "input.docx"
            doc = Document()
            doc.add_paragraph("Tutorial 9")
            doc.add_paragraph("1. Calculate grades")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "≥ 90"
            table.cell(0, 1).text = "A"
            doc.add_paragraph("2. Calculate costs")
            doc.save(path)
            questions = word_document.split_numbered(word_document.read_tutorial(path))
            self.assertEqual(len(questions), 2)
            self.assertIn("≥ 90 | A", questions[0])
            self.assertNotIn("≥ 90", questions[1])

    def test_comments_removed_without_breaking_strings(self):
        source = 'String url = "jdbc:mysql://localhost/db"; // comment\nString s = "/*keep*/"; /* remove */'
        result = ai_answers.remove_comments(source)
        self.assertIn('"jdbc:mysql://localhost/db"', result)
        self.assertIn('"/*keep*/"', result)
        self.assertNotIn("comment", result)
        self.assertNotIn("remove", result)

    def test_path_traversal_rejected(self):
        solution = copy.deepcopy(CONSOLE)
        solution["files"][0]["name"] = "../escape.java"
        with self.assertRaises(ValueError):
            ai_answers.validate_solution(solution)

    def test_capture_without_assertion_rejected(self):
        solution = copy.deepcopy(CONSOLE)
        solution["steps"].pop(0)
        with self.assertRaises(ValueError):
            ai_answers.validate_solution(solution)

    def test_schema_escape_and_multistatement_rejected(self):
        for sql in ["CREATE DATABASE other", "INSERT INTO other.Book (id) VALUES (1)",
                    "CREATE TABLE Book (id INT); DROP DATABASE other"]:
            solution = copy.deepcopy(CONSOLE)
            solution["sql"] = [sql]
            with self.assertRaises(ValueError):
                ai_answers.validate_solution(solution)

    def test_child_does_not_inherit_api_key(self):
        with patch.dict(os.environ, {"OPENAI_API_KEY": "secret", "OTHER_TOKEN": "secret"}):
            self.assertNotIn("OPENAI_API_KEY", java_programs.child_env())
            self.assertNotIn("OTHER_TOKEN", java_programs.child_env())

    def test_cancel_before_process(self):
        event = threading.Event()
        event.set()
        with self.assertRaises(settings.Cancelled):
            java_programs.run_process(["never-run"], settings.ROOT, event)

    def test_invalid_boundaries_rejected(self):
        for starts in [[1, 1], [2, 1], [-1], [3], []]:
            with self.assertRaises(ValueError):
                word_document.questions_from_starts(["title", "one", "two"], starts)

    def test_api_request_and_response(self):
        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                pass

            def read(self):
                return json.dumps({"status": "completed", "output": [
                    {"type": "message", "content": [{"type": "output_text", "text": '{"ok": true}'}]}
                ]}).encode()

        with patch("urllib.request.urlopen", return_value=Response()) as request:
            result = ai_answers.AI("test-key", "test-model", threading.Event()).ask(
                "instruction", "question", ai_answers.object_schema({"ok": {"type": "boolean"}}))
            self.assertTrue(result["ok"])
            body = json.loads(request.call_args.args[0].data)
            self.assertFalse(body["store"])
            self.assertTrue(body["text"]["format"]["strict"])


@unittest.skipUnless(os.environ.get("TUTORIAL_GUI_TESTS") == "1", "Visible desktop integration tests")
class DesktopTests(unittest.TestCase):
    def test_real_swing_console_and_document(self):
        folder = settings.ROOT / "output" / "verification"
        folder.mkdir(parents=True, exist_ok=True)
        results = []
        options = {"db_host": "127.0.0.1", "db_port": "3306", "db_user": "root", "db_password": ""}
        for index, fixture in enumerate([CONSOLE, SWING], 1):
            solution = ai_answers.validate_solution(copy.deepcopy(fixture))
            attempt = folder / ("question_" + str(index) + "_" + uuid.uuid4().hex[:6])
            screenshots = java_programs.compile_and_run(solution, attempt, options, threading.Event(), print)
            self.assertGreaterEqual(len(screenshots), 1)
            results.append({"question": f"{index}. Local integration test", "solution": solution,
                            "screenshots": screenshots})
        path = word_document.write_document(folder / "Verification.docx", "Tutorial verification",
                                  "Sample Student", "TEST001", results)
        doc = Document(path)
        self.assertEqual(len(doc.inline_shapes), sum(len(r["screenshots"]) for r in results))
        self.assertIn("Registration Number: TEST001", [p.text for p in doc.paragraphs])
        print("Verified document:", path)

    def test_bad_assertion_fails_instead_of_exporting(self):
        solution = copy.deepcopy(CONSOLE)
        solution["steps"][0]["value"] = "Sum: 999"
        folder = settings.ROOT / "output" / "verification" / ("failure_" + uuid.uuid4().hex[:6])
        with self.assertRaisesRegex(RuntimeError, "Expected"):
            java_programs.compile_and_run(solution, folder, {}, threading.Event(), print)
        self.assertFalse((folder / "screenshots/captions.tsv").exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
