from pathlib import Path
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from settings import SUBJECT


def read_tutorial(path):
    path = Path(path)
    if path.suffix.lower() != ".docx":
        raise ValueError("Choose a Word .docx file. Save older .doc files as .docx first.")
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 60_000_000:
            raise ValueError("This document is too large. Use a tutorial under 60 MB expanded.")
    doc = Document(path)
    blocks = []
    for element in doc.element.body:
        if element.tag == qn("w:p"):
            text = Paragraph(element, doc).text.strip()
        elif element.tag == qn("w:tbl"):
            table = Table(element, doc)
            text = "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
        else:
            continue
        if text:
            blocks.append(text)
    if not blocks:
        raise ValueError("No readable text found. Scanned/image-only tutorials need OCR first.")
    if sum(map(len, blocks)) > 65_000:
        raise ValueError("Please split tutorials longer than 65,000 characters.")
    return blocks


def split_numbered(blocks):
    pattern = re.compile(r"^\s*(?:(?:question|q)\s*\.?\s*)?(\d{1,2})\s*[.):]\s*", re.I)
    starts = []
    for i, block in enumerate(blocks):
        match = pattern.match(block)
        if match:
            starts.append((i, int(match[1])))
    if not starts or [n for _, n in starts] != list(range(1, len(starts) + 1)):
        return None
    return questions_from_starts(blocks, [i for i, _ in starts])


def questions_from_starts(blocks, starts):
    if not starts or starts != sorted(set(starts)) or starts[0] < 0 or starts[-1] >= len(blocks):
        raise ValueError("Question boundaries could not be verified. Use numbered questions.")
    questions = []
    ends = starts[1:] + [len(blocks)]
    for start, end in zip(starts, ends):
        question = "\n\n".join(blocks[start:end])
        questions.append(question)
    return questions


def safe_filename(text):
    return re.sub(r"[^\w .-]", "_", text).strip(" .")[:90] or "Tutorial"


def add_code(doc, code):
    for line in code.splitlines():
        paragraph = doc.add_paragraph(style="Code")
        paragraph.add_run(line or " ")


def write_document(path, title, name, registration, results, partial=False):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    code = doc.styles.add_style("Code", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(9)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1
    doc.add_paragraph(title, "Title")
    doc.add_paragraph("Name: " + name)
    doc.add_paragraph("Registration Number: " + registration)
    doc.add_paragraph("Subject: " + SUBJECT)
    if partial:
        doc.add_paragraph("Incomplete — some questions could not be verified.")
    for index, result in enumerate(results, 1):
        if index > 1:
            doc.add_page_break()
        doc.add_heading(f"Question {index}", level=1)
        for part in result["question"].split("\n\n"):
            doc.add_paragraph(part)
        if result.get("error"):
            doc.add_paragraph("Answer pending — execution could not be verified.")
            continue
        solution = result["solution"]
        if solution["assumptions"]:
            doc.add_heading("Assumptions", level=2)
            for assumption in solution["assumptions"]:
                doc.add_paragraph(assumption)
        doc.add_heading("Code / Answer", level=2)
        if solution["sql"]:
            doc.add_heading("Database setup", level=3)
            add_code(doc, ";\n".join(s.rstrip("; \n") for s in solution["sql"]) + ";")
        for file in solution["files"]:
            doc.add_heading(file["name"], level=3)
            add_code(doc, file["code"])
        doc.add_heading("Output", level=2)
        for picture, caption in result["screenshots"]:
            paragraph = doc.add_paragraph(caption)
            paragraph.paragraph_format.keep_with_next = True
            doc.add_picture(str(picture), width=Inches(6.1))
    temporary = path.with_suffix(".tmp.docx")
    doc.save(temporary)
    temporary.replace(path)
    return path
